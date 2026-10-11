"""L13 repro 1 — shared review-workbook builder, synthetic rows only (no DB)."""
import zipfile, re
from openpyxl import load_workbook
from engine.exporters.review_workbook import (
    ColumnDef, DecisionColumnDef, FreeTextColumnDef, InstructionsConfig,
    create_review_workbook)

cols = [ColumnDef(key="paper_id", header="Paper ID"),
        ColumnDef(key="title", header="Title"),
        ColumnDef(key="abstract", header="Abstract", truncate=20),
        ColumnDef(key="big", header="Big")]
dec = [DecisionColumnDef(key="PI_decision", header="PI_decision (A/B)", valid_values=["A", "B"])]
ft = [FreeTextColumnDef(key="PI_notes", header="PI_notes (optional)")]
rows = [{"paper_id": 7, "title": "T one", "abstract": "x" * 50, "big": "y" * 40000},
        {"paper_id": 8, "title": "", "abstract": None, "big": ""}]
import io, contextlib
with contextlib.redirect_stdout(io.StringIO()):
    create_review_workbook("q.xlsx", rows, cols, dec,
                           InstructionsConfig(review_name="synth", row_count=999),
                           free_text_columns=ft)
z = zipfile.ZipFile("q.xlsx")
sheet = z.read("xl/worksheets/sheet2.xml").decode()
print("[a] dataValidation element:", re.search(r"<dataValidation [^>]*>", sheet).group(0))
wb = load_workbook("q.xlsx")
ws = wb["Review Queue"]
print("[b] truncate=20 cell:", repr(ws.cell(2, 3).value), "(no marker)")
print("[c] 40000-char value stored length:", len(ws.cell(2, 4).value))
print("[d] row with empty title, as reloaded:", [c.value for c in ws[3]])
print("[e] Instructions 'Total rows' cell vs real rows:",
      [r[1].value for r in wb["Instructions"].iter_rows() if r[0].value == "Total rows in Review Queue"], "real", len(rows))

try:
    with contextlib.redirect_stdout(io.StringIO()):
        create_review_workbook("empty.xlsx", [], cols, dec, InstructionsConfig(review_name="synth"))
    print("[f] zero rows: ok")
except Exception as e:
    import os
    print("[f] zero rows ->", type(e).__name__, e, "| file written:", os.path.exists("empty.xlsx"))

# the module docstring's own usage example (dict column defs)
try:
    with contextlib.redirect_stdout(io.StringIO()):
        create_review_workbook("doc.xlsx", [{"paper_id": 1, "title": "t"}],
            columns=[{"key": "paper_id", "header": "Paper ID", "width": 10}],
            decision_columns=[{"key": "PI_decision", "header": "PI_decision", "valid_values": ["A"], "width": 25}],
            instructions=InstructionsConfig(review_name="synth"))
    print("[g] docstring usage: ok")
except Exception as e:
    print("[g] docstring usage example ->", type(e).__name__, e)
