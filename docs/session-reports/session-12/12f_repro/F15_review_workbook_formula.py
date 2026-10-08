"""F15 — the repo's shared adjudication workbook builder, fed a formula-leading title/rationale.
Writes only into cwd (~/scratch/12f/g3). No db, no model, no network."""
import contextlib, io
from openpyxl import load_workbook
from engine.exporters.review_workbook import (
    ColumnDef, DecisionColumnDef, InstructionsConfig, create_review_workbook)
rows = [{"paper_id": 1, "title": '=HYPERLINK("http://example.invalid","Robotic suturing")',
         "primary_rationale": "=1+1"},
        {"paper_id": 2, "title": "-12% leak rate after robotic anastomosis", "primary_rationale": "@note"}]
cols = [ColumnDef(key="paper_id", header="Paper ID"), ColumnDef(key="title", header="Title"),
        ColumnDef(key="primary_rationale", header="Primary Rationale")]
dec = [DecisionColumnDef(key="PI_decision", header="PI_decision", valid_values=["INCLUDE", "EXCLUDE"])]
with contextlib.redirect_stdout(io.StringIO()):
    create_review_workbook("f15_queue.xlsx", rows, cols, dec, InstructionsConfig(review_name="synth"))
ws = load_workbook("f15_queue.xlsx")["Review Queue"]
for r in ws.iter_rows(min_row=2, max_col=3):
    print([(c.data_type, c.value) for c in r])
rows[0]["title"] = "Title with a form feed \x0c from parsed text"
try:
    with contextlib.redirect_stdout(io.StringIO()):
        create_review_workbook("f15_queue2.xlsx", rows, cols, dec, InstructionsConfig(review_name="synth"))
    print("control char: written")
except Exception as e:
    print("control char ->", type(e).__name__)
